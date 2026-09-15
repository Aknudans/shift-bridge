import sys
import time
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import sync_playwright, Page
import pandas as pd

# Direcciones y selectores del sitio que usan tanto la fase de comparación
# como la de creación/edición. Si el sitio cambia, se ajustan acá.

BASE_URL = "https://externoslof.shiftlabor.com/Funcionalidades/Externos/ProveedorTrabajador.aspx"
# Va por IP y no por "localhost" porque en Windows eso apunta a IPv6 y Chrome
# escucha en IPv4.
CDP_URL = "http://127.0.0.1:9222"

SELECTOR_FILTRO_RUT = "#grillaExternosProveedorTrabajadores_DXFREditorcol2_I"

# La grilla de trabajadores. Se busca dentro de la tabla y no por clase de
# fila, porque una fila abierta antes cambia de clase y se dejaría de encontrar.
SELECTOR_TABLA_GRILLA = "#grillaExternosProveedorTrabajadores_DXMainTable"

# Combo de arriba a la izquierda donde se elige el grupo proveedor.
SELECTOR_INPUT_GRUPO_PROVEEDOR = "#MJJerarquia00_I"

# Menú lateral, para llegar a Trabajadores paso a paso.
SELECTOR_BOTON_MENU = "#icono_abrir_menu"
SELECTOR_MENU_EXTERNOS = "#mf_cab_01"
SELECTOR_MENU_TRABAJADORES = "#mf_cab_ll_01_01"

# El Excel trae proveedor, cargo y tienda con este prefijo adelante, pero el
# sitio no siempre lo muestra. Se saca de ambos lados antes de comparar.
PREFIJO_CATALOGO = "LOGISTICA FALABELLA/"


# Comparar texto del Excel contra texto del sitio sin que estorben las
# mayúsculas, los espacios de más ni el prefijo del catálogo.

def normalizar_texto(valor: Optional[str]) -> str:
    if valor is None:
        return ""
    return str(valor).strip().upper()


def quitar_prefijo_catalogo(valor: Optional[str]) -> str:
    texto = normalizar_texto(valor)
    if texto.startswith(normalizar_texto(PREFIJO_CATALOGO)):
        return texto[len(PREFIJO_CATALOGO):].strip()
    return texto


# El Excel llega con los encabezados de la planilla original para que el
# usuario pegue el bloque tal cual. Acá se traducen a los nombres que usa el
# resto del código.
RENOMBRE_COLUMNAS_ENTRADA = {
    "rut": "RUT",
    "nombre": "NOMBRES",
    "sexo": "SEXO",
    "cargo": "CARGO",
    "desde": "fechaContratacion",
    "hasta": "fechaTermino",
    "afp": "AFP",
    "isapre": "ISAPRE",
}


# La planilla escribe el sexo como M o F y el sitio espera la palabra completa.
def normalizar_sexo(valor: Optional[str]) -> str:
    if valor is None:
        return ""
    t = str(valor).strip().upper()
    if t in ("M", "MASCULINO"):
        return "Masculino"
    if t in ("F", "FEMENINO"):
        return "Femenino"
    if t in ("", "NAN", "NONE"):
        return ""
    return str(valor).strip()


# Abre el Excel de entrada, lo deja con los nombres de columna internos y
# corta el programa si falta alguna columna obligatoria.
def cargar_excel(path: str, columnas_requeridas: list[str]) -> pd.DataFrame:
    df = pd.read_excel(path, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    df = df.rename(columns=RENOMBRE_COLUMNAS_ENTRADA)

    if "SEXO" in df.columns:
        df["SEXO"] = df["SEXO"].map(normalizar_sexo)

    faltantes = [c for c in columnas_requeridas if c not in df.columns]
    if faltantes:
        print(f"ERROR: Faltan columnas obligatorias en el Excel de entrada: {faltantes}")
        sys.exit(1)

    return df


# El bot no abre Chrome ni maneja claves: se cuelga de la ventana que la
# persona ya dejó abierta con la sesión iniciada.
def conectar_a_chrome_existente():
    playwright = sync_playwright().start()
    try:
        browser = playwright.chromium.connect_over_cdp(CDP_URL)
    except Exception as e:
        print("ERROR: No se pudo conectar a Chrome en el puerto 9222.")
        print("Verificar que Chrome esté abierto con --remote-debugging-port=9222 y que esa ventana siga abierta.")
        print(f"Detalle técnico: {e}")
        sys.exit(1)

    context = browser.contexts[0] if browser.contexts else browser.new_context()
    page = context.pages[0] if context.pages else context.new_page()
    return playwright, browser, page


# Las dos formas de pararse en la lista de trabajadores: ir directo por URL, o
# dar la vuelta por el menú lateral. Lo segundo se usa después de guardar,
# para que el sitio quede en un estado limpio antes de la fila siguiente
# (ojo: eso borra el grupo proveedor elegido y hay que volver a elegirlo).

def asegurar_pagina_trabajadores(page: Page):
    if BASE_URL not in page.url:
        page.goto(BASE_URL)
    page.wait_for_load_state("networkidle")


def navegar_a_trabajadores_por_menu(page: Page):
    page.locator(SELECTOR_BOTON_MENU).click(timeout=5000)
    time.sleep(0.3)
    page.locator(SELECTOR_MENU_EXTERNOS).click(timeout=5000)
    time.sleep(0.3)
    page.locator(SELECTOR_MENU_TRABAJADORES).click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)


# La grilla solo muestra la gente del grupo elegido, así que hay que elegirlo
# antes de buscar cualquier RUT. En el combo el nombre aparece sin el prefijo
# del catálogo, por eso se usa solo lo que viene después de la barra.
def seleccionar_grupo_proveedor(page: Page, nombre_proveedor: str):
    texto_busqueda = nombre_proveedor.split("/")[-1].strip()

    page.locator(SELECTOR_INPUT_GRUPO_PROVEEDOR).click(timeout=5000)
    time.sleep(0.3)
    opcion = page.locator(f"text={texto_busqueda}").first
    opcion.click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)


# Filtra la grilla por RUT y dice si apareció alguien. La espera extra es
# porque el sitio termina de cargar la página antes de terminar de dibujar la
# fila, y sin eso se daba por no encontrada a gente que sí existe.
def buscar_rut(page: Page, rut: str) -> bool:
    filtro = page.locator(SELECTOR_FILTRO_RUT)
    filtro.click()
    filtro.fill("")
    filtro.fill(rut)
    filtro.press("Enter")
    page.wait_for_load_state("networkidle")
    selector_resultado = f"{SELECTOR_TABLA_GRILLA} td:has-text('{rut}')"
    try:
        page.wait_for_selector(selector_resultado, timeout=5000)
    except Exception:
        pass
    time.sleep(0.3)

    return page.locator(selector_resultado).count() > 0


# Deja la grilla sin filtro para la búsqueda siguiente.
def limpiar_filtro(page: Page):
    filtro = page.locator(SELECTOR_FILTRO_RUT)
    filtro.click()
    filtro.fill("")
    filtro.press("Enter")
    page.wait_for_load_state("networkidle")
    time.sleep(0.3)


# Arma el Excel final, una fila por persona y con el color según cómo salió.
# Sirve para las dos fases: solo pide objetos con rut, nombre, estado y
# detalle. Si el archivo está abierto en Excel, guarda con otro nombre en vez
# de perder el trabajo de toda la corrida.
def escribir_reporte(resultados: list, output_path: str, colores_estado: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"

    encabezados = ["RUT", "Nombre (Excel)", "Estado", "Detalle"]
    ws.append(encabezados)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in resultados:
        ws.append([r.rut, r.nombre_excel, r.estado, r.detalle])
        fill = PatternFill(start_color=colores_estado.get(r.estado, "FFFFFF"),
                            end_color=colores_estado.get(r.estado, "FFFFFF"),
                            fill_type="solid")
        for cell in ws[ws.max_row]:
            cell.fill = fill

    for col_cells in ws.columns:
        largo = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(largo + 4, 80)

    try:
        wb.save(output_path)
    except PermissionError:
        alterno = f"{output_path.rsplit('.', 1)[0]}_{int(time.time())}.xlsx"
        print(f"\nADVERTENCIA: no se pudo guardar en '{output_path}' "
              f"(¿está abierto en Excel u otro programa?). Guardando como '{alterno}' en su lugar.")
        wb.save(alterno)
        output_path = alterno

    print(f"\nReporte guardado en: {output_path}")
