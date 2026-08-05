"""
Bot de creación y edición de colaboradores — ShiftLaboral (Fase 2)
====================================================================

QUÉ HACE:
Para cada colaborador de tu Excel de entrada, busca su RUT en ShiftLaboral:

    - Si NO existe -> hace clic en "Crear" y llena el formulario completo con
      los datos del Excel.
    - Si SÍ existe -> hace clic en "editar", lee los valores ACTUALES del
      formulario, calcula qué campos difieren del Excel, y modifica SOLO esos
      campos (nunca toca los que ya coinciden).

En ambos casos, si el Cargo del Excel no calza EXACTO con una opción real del
catálogo (dropdown "Categoría Trabajador"), esa fila se OMITE COMPLETA (no se
crea ni se edita nada de esa persona) y se reporta como error.

Al final genera un Excel de reporte con el resultado de cada fila: CREADO,
EDITADO (con el detalle de qué campos cambiaron), SIN_CAMBIOS, OMITIDO_CARGO
o ERROR.

⚠️ A DIFERENCIA DE buscar_y_comparar.py (Fase 1, solo lectura), este script
   SÍ modifica datos reales en ShiftLaboral. Probarlo primero con 2-3 filas
   de prueba antes de correrlo contra el Excel completo de producción — igual
   protocolo que el descrito en CLAUDE.md sección 8, pero aplicado a Fase 2.

REQUISITOS PREVIOS: los mismos que buscar_y_comparar.py — ver ese archivo o
README.md. Usar preferentemente ejecutar_crear.bat, que automatiza la
apertura de Chrome y el login.

EXCEL DE ENTRADA ESPERADO: mismas columnas que buscar_y_comparar.py, más
fechaContratacion y fechaTermino, que aquí SÍ son obligatorias (se usan para
Inicio Contrato / Fin Contrato al crear o editar).

USO:
    python crear_o_editar.py --input colaboradores.xlsx --output reporte_crear.xlsx

NOTA IMPORTANTE: los selectores de "editar"/"Crear"/"Guardar" fueron
explorados en vivo el 05/08/2026 (ver CLAUDE.md sección 12), pero el botón
"Guardar" nunca se probó realmente durante esa exploración (a propósito, para
no modificar datos sin autorización). Validar con cuidado la primera corrida.
"""

import argparse
import sys
import time
from dataclasses import dataclass, field
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import sync_playwright, Page

# ---------------------------------------------------------------------------
# CONFIGURACIÓN — ajustar aquí si algo cambia en el sitio
# ---------------------------------------------------------------------------

BASE_URL = "https://externoslof.shiftlabor.com/Funcionalidades/Externos/ProveedorTrabajador.aspx"
CDP_URL = "http://localhost:9222"

SELECTOR_FILTRO_RUT = "#grillaExternosProveedorTrabajadores_DXFREditorcol2_I"
SELECTOR_BOTON_EDITAR = 'a[title="editar"]'
SELECTOR_TABLA_GRILLA = "#grillaExternosProveedorTrabajadores_DXMainTable"
SELECTOR_INPUT_GRUPO_PROVEEDOR = "#MJJerarquia00_I"
SELECTOR_BOTON_CREAR_NUEVO = "text=Crear"
# ⚠️ El input real editable es "#txtRutVer_I" — "#txtRutVer_Raw" es un input
# oculto (type="hidden") que NO hay que tocar directamente (confirmado en vivo).
SELECTOR_RUT_CREAR = "#txtRutVer_I"
SELECTOR_BTN_ACEPTAR_RUT = "#btnAceptaRut"

# Botones del formulario (editar/crear). CONFIRMADOS EN VIVO 05/08/2026 por
# inspección de DOM — NUNCA se hizo clic real en #btnGuardar durante la
# exploración, para no modificar datos reales sin autorización. Validar con
# cuidado la primera corrida (ver CLAUDE.md sección 12).
SELECTOR_BTN_GUARDAR = "#btnGuardar"
SELECTOR_BTN_CANCELAR = "#btnCancelar"

# Fragmento estable del id del botón ">" que confirma el Proveedor elegido y
# habilita las secciones de Categoría Trabajador (Cargo) y Tiendas.
ID_FRAGMENTO_CONFIRMAR_PROVEEDOR = "btnCargarClientesCategoria"

# Selectores del menú lateral (ver buscar_y_comparar.py / CLAUDE.md sección 6)
# — usados para RESETEAR la vista de Trabajadores después de cada Crear/Editar
# exitoso, por pedido explícito del usuario: en vez de solo recargar la URL,
# se navega pasando por el menú (hamburguesa -> Externos -> Trabajadores)
# para asegurar que el servidor quede en un estado limpio antes de la
# siguiente fila.
SELECTOR_BOTON_MENU = "#icono_abrir_menu"
SELECTOR_MENU_EXTERNOS = "#mf_cab_01"
SELECTOR_MENU_TRABAJADORES = "#mf_cab_ll_01_01"

# Campos de texto simples: mismo ID en "editar" y en "Crear" (gran ventaja,
# confirmado en vivo). Mapeo etiqueta -> id.
CAMPOS_TEXTO = {
    "Nombres:": "txtNombres",
    "Apellido Paterno:": "txtApellidoPaterno",
    "Apellido Materno:": "txtApellidoMaterno",
    "Sueldo Base:": "txtSueldoBase",
}

# Combobox de selección única (sin checkboxes) — fragmento estable del id,
# terminado en "_I" para el input visible (evitar "_VI", que es el valor
# oculto). Mismo fragmento en editar y crear, solo cambia el prefijo de fila.
COMBOBOX_SIMPLES = {
    "Sexo:": "cbSexo",
    "AFP:": "cbAFP",
    "Sistema de Salud:": "cbIsapre",
}

SELECTOR_FECHA_INICIO = "#calendarioFechaInicio_txtCalendar"
SELECTOR_FECHA_TERMINO = "#calendarioFechaTermino_txtCalendar"

COLUMNAS_EXCEL_REQUERIDAS = [
    "RUT", "NOMBRES", "apellidoPaterno", "apellidoMaterno", "SEXO", "AFP", "ISAPRE",
    "sueldoBase", "PROVEEDOR", "CARGO", "TIENDA", "fechaContratacion", "fechaTermino",
]

PREFIJO_CATALOGO = "LOGISTICA FALABELLA/"


# ---------------------------------------------------------------------------
# MODELOS DE DATOS
# ---------------------------------------------------------------------------

@dataclass
class ResultadoFila:
    rut: str
    nombre_excel: str
    estado: str  # "CREADO" | "EDITADO" | "SIN_CAMBIOS" | "OMITIDO_CARGO" | "ERROR"
    detalle: str = ""


# ---------------------------------------------------------------------------
# UTILIDADES
# ---------------------------------------------------------------------------

def normalizar_texto(valor: Optional[str]) -> str:
    if valor is None:
        return ""
    return str(valor).strip().upper()


def quitar_prefijo_catalogo(valor: Optional[str]) -> str:
    texto = normalizar_texto(valor)
    if texto.startswith(normalizar_texto(PREFIJO_CATALOGO)):
        return texto[len(PREFIJO_CATALOGO):].strip()
    return texto


def formatear_fecha(valor: Optional[str]) -> str:
    """Convierte una fecha del Excel (ej. '2026-04-08 00:00:00') al formato
    dd/mm/aaaa que usa el formulario de ShiftLaboral."""
    if valor is None or str(valor).strip() == "" or str(valor).lower() == "nan":
        return ""
    fecha = pd.to_datetime(valor)
    return fecha.strftime("%d/%m/%Y")


def cargar_excel(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, dtype=str)
    df.columns = [c.strip() for c in df.columns]

    faltantes = [c for c in COLUMNAS_EXCEL_REQUERIDAS if c not in df.columns]
    if faltantes:
        print(f"ERROR: Faltan columnas obligatorias en el Excel de entrada: {faltantes}")
        sys.exit(1)

    return df


# ---------------------------------------------------------------------------
# NAVEGACIÓN Y BÚSQUEDA (comparte lógica con buscar_y_comparar.py)
# ---------------------------------------------------------------------------

def conectar_a_chrome_existente():
    playwright = sync_playwright().start()
    try:
        browser = playwright.chromium.connect_over_cdp(CDP_URL)
    except Exception as e:
        print("ERROR: No se pudo conectar a Chrome en el puerto 9222.")
        print("¿Abriste Chrome con --remote-debugging-port=9222 y dejaste esa ventana abierta?")
        print(f"Detalle técnico: {e}")
        sys.exit(1)

    context = browser.contexts[0] if browser.contexts else browser.new_context()
    page = context.pages[0] if context.pages else context.new_page()
    return playwright, browser, page


def asegurar_pagina_trabajadores(page: Page):
    if BASE_URL not in page.url:
        page.goto(BASE_URL)
    page.wait_for_load_state("networkidle")


def seleccionar_grupo_proveedor(page: Page, nombre_proveedor: str):
    """Selecciona el grupo/proveedor correcto en el combo superior. Ver
    CLAUDE.md sección 6 para el detalle de por qué esto está automatizado."""
    texto_busqueda = nombre_proveedor.split("/")[-1].strip()

    page.locator(SELECTOR_INPUT_GRUPO_PROVEEDOR).click(timeout=5000)
    time.sleep(0.3)
    opcion = page.locator(f"text={texto_busqueda}").first
    opcion.click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)


def resetear_vista_trabajadores(page: Page):
    """Vuelve a la vista de Trabajadores navegando por el menú lateral
    (hamburguesa -> Externos -> Trabajadores), en vez de solo recargar la
    URL. Se llama después de cada Crear/Editar exitoso (pedido explícito del
    usuario) para asegurar un estado de servidor limpio antes de la siguiente
    fila, evitando arrastrar cualquier estado residual del guardado anterior.

    Como esto navega a una página nueva, el Grupo Proveedor seleccionado se
    pierde — quien llama debe volver a seleccionarlo para la siguiente fila.
    """
    page.locator(SELECTOR_BOTON_MENU).click(timeout=5000)
    time.sleep(0.3)
    page.locator(SELECTOR_MENU_EXTERNOS).click(timeout=5000)
    time.sleep(0.3)
    page.locator(SELECTOR_MENU_TRABAJADORES).click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)


def buscar_rut(page: Page, rut: str) -> bool:
    """Ver buscar_y_comparar.py para el detalle del fix de race condition."""
    filtro = page.locator(SELECTOR_FILTRO_RUT)
    filtro.click()
    filtro.fill("")
    filtro.fill(rut)
    filtro.press("Enter")
    page.wait_for_load_state("networkidle")
    selector_resultado = f"{SELECTOR_TABLA_GRILLA} td:has-text('{rut}')"
    try:
        page.wait_for_selector(selector_resultado, timeout=5000)
    except Exception:
        pass
    time.sleep(0.3)

    return page.locator(selector_resultado).count() > 0


def limpiar_filtro(page: Page):
    filtro = page.locator(SELECTOR_FILTRO_RUT)
    filtro.click()
    filtro.fill("")
    filtro.press("Enter")
    page.wait_for_load_state("networkidle")
    time.sleep(0.3)


def cerrar_formulario(page: Page):
    """Cierra cualquier vista/formulario abierto (editar, ver o crear) con el
    botón Cancelar, sin guardar. CRÍTICO no omitir esto entre filas — ver
    CLAUDE.md sección 6 Paso 11 (bug de fila "pegada" en modo edición)."""
    try:
        page.locator(SELECTOR_BTN_CANCELAR).click(timeout=3000)
        page.wait_for_load_state("networkidle")
        time.sleep(0.3)
    except Exception:
        pass


# ---------------------------------------------------------------------------
# LECTURA Y ESCRITURA DE CAMPOS DEL FORMULARIO (editar y crear comparten IDs)
# ---------------------------------------------------------------------------

def _input_por_etiqueta(page: Page, etiqueta: str) -> Optional[str]:
    """Devuelve el id del input/select asociado a una etiqueta visible.

    CONFIRMADO EN VIVO 05/08/2026: a diferencia de Nombres/Apellidos/AFP (con
    etiqueta y caja en la misma fila de tabla, td + td), las etiquetas
    "Proveedores:", "Categoría Trabajador:" y "Tiendas:" están en su propia
    línea, con la caja completa DEBAJO (no al lado) — "closest('td') +
    nextElementSibling" no encuentra nada ahí y siempre devuelve null. En vez
    de asumir estructura de tabla, se toma el primer input/select/textarea
    VISIBLE que aparece después de la etiqueta en el orden del documento.
    """
    return page.evaluate(
        """(etiqueta) => {
            const el = Array.from(document.querySelectorAll('*')).find(e =>
                e.offsetParent !== null && e.children.length === 0 && e.textContent.trim() === etiqueta
            );
            if (!el) return null;
            const candidatos = Array.from(document.querySelectorAll('input, select, textarea'))
                .filter(i => i.offsetParent !== null);
            for (const input of candidatos) {
                if (el.compareDocumentPosition(input) & Node.DOCUMENT_POSITION_FOLLOWING) {
                    return input.id;
                }
            }
            return null;
        }""",
        etiqueta,
    )


def leer_campo_texto(page: Page, campo_id: str) -> str:
    try:
        return page.locator(f"#{campo_id}").input_value(timeout=3000).strip()
    except Exception:
        return ""


def escribir_campo_texto(page: Page, campo_id: str, valor: str):
    campo = page.locator(f"#{campo_id}")
    campo.click()
    campo.fill("")
    campo.fill(valor)
    campo.press("Tab")


def leer_combobox_simple(page: Page, fragmento: str) -> str:
    try:
        return page.locator(f'[id*="{fragmento}"][id$="_I"]').input_value(timeout=3000).strip()
    except Exception:
        return ""


def escribir_combobox_simple(page: Page, fragmento: str, valor_deseado: str):
    """Abre el combobox de selección única y hace clic en la opción cuyo
    texto calza EXACTO (ignorando mayúsculas/espacios) con valor_deseado."""
    input_visible = page.locator(f'[id*="{fragmento}"][id$="_I"]')
    input_visible.click(timeout=5000)
    time.sleep(0.3)
    opcion = page.locator(f"text={valor_deseado}").first
    opcion.click(timeout=5000)
    time.sleep(0.3)


def leer_multiselect(page: Page, etiqueta: str) -> list[str]:
    """Lee el valor actual (semicolon-separated en el input visible cuando
    hay más de uno) de un combobox multi-select (Proveedores/Cargo/Tiendas)."""
    campo_id = _input_por_etiqueta(page, etiqueta)
    if not campo_id:
        return []
    try:
        valor = page.locator(f"#{campo_id}").input_value(timeout=3000).strip()
    except Exception:
        return []
    if not valor:
        return []
    return [v.strip() for v in valor.split(";") if v.strip()]


def obtener_opciones_multiselect(page: Page, etiqueta: str) -> list[dict]:
    """Abre el dropdown de un multi-select y devuelve todas sus opciones
    visibles con su texto y si están marcadas. NO cierra el dropdown.

    CONFIRMADO EN VIVO 05/08/2026: cada opción del checkbox list de DevExpress
    se renderiza como DOS elementos `.dxeListBoxItem` SEPARADOS y consecutivos
    — uno con el `<input type="checkbox">` (clase "dxeC", sin texto) y el
    siguiente con el texto de la etiqueta (clase "dxeT", sin checkbox). No son
    un único elemento con ambos. Hay que emparejarlos por posición.
    """
    campo_id = _input_por_etiqueta(page, etiqueta)
    if not campo_id:
        return []
    page.locator(f"#{campo_id}").click(timeout=5000)
    time.sleep(0.3)
    return page.evaluate(
        """() => {
            const items = Array.from(document.querySelectorAll('.dxeListBoxItem'))
                .filter(e => e.offsetParent !== null);
            const opciones = [];
            for (let i = 0; i < items.length - 1; i++) {
                const chk = items[i].querySelector('input[type="checkbox"]');
                if (!chk) continue;
                const texto = items[i + 1].textContent.trim();
                if (!texto || texto === 'Seleccionar Todos') continue;
                opciones.push({ texto, marcado: chk.checked });
            }
            return opciones;
        }"""
    )


def establecer_multiselect_valor_unico(page: Page, etiqueta: str, valor_deseado: str) -> bool:
    """Deja UN ÚNICO valor marcado en un combobox multi-select (Proveedores/
    Cargo/Tiendas): desmarca cualquier opción marcada que no sea la deseada,
    y marca la deseada si no lo estaba. Devuelve False si valor_deseado no
    existe EXACTO entre las opciones reales (candidato a fila omitida).

    Todo el emparejamiento checkbox+texto y el clic se hacen en un solo
    page.evaluate (necesita la referencia viva al <input>, no solo su texto)
    — ver nota de obtener_opciones_multiselect sobre por qué están separados.
    """
    campo_id = _input_por_etiqueta(page, etiqueta)
    if not campo_id:
        return False
    page.locator(f"#{campo_id}").click(timeout=5000)
    time.sleep(0.3)

    resultado = page.evaluate(
        """(valorDeseado) => {
            const norm = (s) => (s || '').trim().toUpperCase();
            const items = Array.from(document.querySelectorAll('.dxeListBoxItem'))
                .filter(e => e.offsetParent !== null);
            const pares = [];
            for (let i = 0; i < items.length - 1; i++) {
                const chk = items[i].querySelector('input[type="checkbox"]');
                if (!chk) continue;
                const texto = items[i + 1].textContent.trim();
                if (!texto || texto === 'Seleccionar Todos') continue;
                pares.push({ checkbox: chk, texto, marcado: chk.checked });
            }
            const deseadoNorm = norm(valorDeseado);
            const existe = pares.some(p => norm(p.texto) === deseadoNorm);
            if (!existe) return { existe: false };
            for (const p of pares) {
                const coincide = norm(p.texto) === deseadoNorm;
                if (p.marcado !== coincide) {
                    p.checkbox.click();
                }
            }
            return { existe: true };
        }""",
        valor_deseado,
    )

    page.keyboard.press("Escape")
    time.sleep(0.2)
    return resultado["existe"]


def validar_cargo_existe(page: Page, cargo_deseado: str) -> bool:
    """Solo valida (sin modificar nada) si cargo_deseado existe EXACTO en el
    catálogo real de Categoría Trabajador. Cierra el dropdown al terminar."""
    opciones = obtener_opciones_multiselect(page, "Categoría Trabajador:")
    texto_normalizado = normalizar_texto(cargo_deseado)
    existe = any(normalizar_texto(o["texto"]) == texto_normalizado for o in opciones)
    page.keyboard.press("Escape")
    time.sleep(0.2)
    return existe


def confirmar_proveedor_seleccionado(page: Page):
    """Hace clic en el botón '>' que confirma el Proveedor y habilita las
    secciones de Categoría Trabajador (Cargo) y Tiendas."""
    page.locator(f'[id*="{ID_FRAGMENTO_CONFIRMAR_PROVEEDOR}"]').first.click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)


# ---------------------------------------------------------------------------
# LECTURA DE VALORES ACTUALES (para decidir qué difiere antes de editar)
# ---------------------------------------------------------------------------

def leer_valores_actuales(page: Page) -> dict:
    """Lee todos los valores actuales del formulario de edición ya abierto."""
    datos = {}
    for etiqueta, campo_id in CAMPOS_TEXTO.items():
        datos[campo_id] = leer_campo_texto(page, campo_id)
    for etiqueta, fragmento in COMBOBOX_SIMPLES.items():
        datos[fragmento] = leer_combobox_simple(page, fragmento)
    datos["proveedores"] = leer_multiselect(page, "Proveedores:")
    datos["cargo"] = leer_multiselect(page, "Categoría Trabajador:")
    datos["tienda"] = leer_multiselect(page, "Tiendas:")
    try:
        datos["fecha_inicio"] = page.locator(SELECTOR_FECHA_INICIO).input_value(timeout=3000).strip()
    except Exception:
        datos["fecha_inicio"] = ""
    try:
        datos["fecha_termino"] = page.locator(SELECTOR_FECHA_TERMINO).input_value(timeout=3000).strip()
    except Exception:
        datos["fecha_termino"] = ""
    return datos


# ---------------------------------------------------------------------------
# FLUJO: EDITAR TRABAJADOR EXISTENTE
# ---------------------------------------------------------------------------

def editar_trabajador_existente(page: Page, fila: pd.Series) -> ResultadoFila:
    rut = str(fila["RUT"]).strip()
    nombre_completo = f"{fila['NOMBRES']} {fila['apellidoPaterno']}"

    page.locator(SELECTOR_BOTON_EDITAR).first.click(timeout=5000)
    page.wait_for_load_state("networkidle")
    try:
        page.wait_for_selector("#txtNombres", timeout=5000)
    except Exception:
        pass
    time.sleep(0.3)

    # 1) Validar Cargo ANTES de tocar cualquier campo. Si no calza, se omite
    #    la fila completa (decisión de negocio confirmada, CLAUDE.md 12.4).
    cargo_deseado = fila["CARGO"]
    if not validar_cargo_existe(page, cargo_deseado):
        cerrar_formulario(page)
        return ResultadoFila(
            rut=rut, nombre_excel=nombre_completo, estado="OMITIDO_CARGO",
            detalle=f"Cargo '{cargo_deseado}' no existe en el catálogo real. Fila completa omitida.",
        )

    # 2) Leer valores actuales y calcular diffs.
    actuales = leer_valores_actuales(page)
    campos_cambiados = []

    def _tal_vez_actualizar_texto(etiqueta, campo_id, valor_excel):
        if normalizar_texto(actuales.get(campo_id)) != normalizar_texto(valor_excel):
            escribir_campo_texto(page, campo_id, str(valor_excel))
            campos_cambiados.append(etiqueta)

    _tal_vez_actualizar_texto("Nombres", "txtNombres", fila["NOMBRES"])
    _tal_vez_actualizar_texto("Apellido Paterno", "txtApellidoPaterno", fila["apellidoPaterno"])
    _tal_vez_actualizar_texto("Apellido Materno", "txtApellidoMaterno", fila["apellidoMaterno"])
    _tal_vez_actualizar_texto("Sueldo Base", "txtSueldoBase", fila["sueldoBase"])

    if normalizar_texto(actuales.get("cbSexo")) != normalizar_texto(fila["SEXO"]):
        escribir_combobox_simple(page, "cbSexo", fila["SEXO"])
        campos_cambiados.append("Sexo")
    if normalizar_texto(actuales.get("cbAFP")) != normalizar_texto(fila["AFP"]):
        escribir_combobox_simple(page, "cbAFP", fila["AFP"])
        campos_cambiados.append("AFP")
    if normalizar_texto(actuales.get("cbIsapre")) != normalizar_texto(fila["ISAPRE"]):
        escribir_combobox_simple(page, "cbIsapre", fila["ISAPRE"])
        campos_cambiados.append("Sistema de Salud")

    fecha_inicio_excel = formatear_fecha(fila["fechaContratacion"])
    fecha_termino_excel = formatear_fecha(fila["fechaTermino"])
    if actuales.get("fecha_inicio") != fecha_inicio_excel:
        escribir_campo_texto(page, "calendarioFechaInicio_txtCalendar", fecha_inicio_excel)
        campos_cambiados.append("Inicio Contrato")
    if actuales.get("fecha_termino") != fecha_termino_excel:
        escribir_campo_texto(page, "calendarioFechaTermino_txtCalendar", fecha_termino_excel)
        campos_cambiados.append("Fin Contrato")

    # NOTA: a diferencia de la vista de solo lectura (Fase 1), en el formulario
    # de editar/crear los 3 catálogos (Proveedores, Categoría Trabajador,
    # Tiendas) SÍ muestran el prefijo "LOGISTICA FALABELLA/" — confirmado en
    # vivo. Por eso aquí se pasa el valor COMPLETO del Excel (con prefijo) a
    # establecer_multiselect_valor_unico, sin recortarlo.
    proveedor_excel = quitar_prefijo_catalogo(fila["PROVEEDOR"])
    proveedores_actuales = [quitar_prefijo_catalogo(v) for v in actuales.get("proveedores", [])]
    if proveedores_actuales != [proveedor_excel]:
        if not establecer_multiselect_valor_unico(page, "Proveedores:", fila["PROVEEDOR"]):
            cerrar_formulario(page)
            return ResultadoFila(
                rut=rut, nombre_excel=nombre_completo, estado="ERROR",
                detalle=f"Proveedor '{fila['PROVEEDOR']}' no existe en el catálogo real.",
            )
        campos_cambiados.append("Proveedores")

    cargo_actual = [quitar_prefijo_catalogo(v) for v in actuales.get("cargo", [])]
    if cargo_actual != [quitar_prefijo_catalogo(cargo_deseado)]:
        establecer_multiselect_valor_unico(page, "Categoría Trabajador:", cargo_deseado)
        campos_cambiados.append("Cargo")

    tienda_excel = quitar_prefijo_catalogo(fila["TIENDA"])
    tienda_actual = [quitar_prefijo_catalogo(v) for v in actuales.get("tienda", [])]
    if tienda_actual != [tienda_excel]:
        if not establecer_multiselect_valor_unico(page, "Tiendas:", fila["TIENDA"]):
            cerrar_formulario(page)
            return ResultadoFila(
                rut=rut, nombre_excel=nombre_completo, estado="ERROR",
                detalle=f"Tienda '{fila['TIENDA']}' no existe en el catálogo real.",
            )
        campos_cambiados.append("Tiendas")

    if not campos_cambiados:
        cerrar_formulario(page)
        return ResultadoFila(rut=rut, nombre_excel=nombre_completo, estado="SIN_CAMBIOS",
                              detalle="Todos los datos ya coincidían.")

    page.locator(SELECTOR_BTN_GUARDAR).click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)

    return ResultadoFila(rut=rut, nombre_excel=nombre_completo, estado="EDITADO",
                          detalle="Campos actualizados: " + ", ".join(campos_cambiados))


# ---------------------------------------------------------------------------
# FLUJO: CREAR TRABAJADOR NUEVO
# ---------------------------------------------------------------------------

def crear_trabajador_nuevo(page: Page, fila: pd.Series) -> ResultadoFila:
    rut = str(fila["RUT"]).strip()
    nombre_completo = f"{fila['NOMBRES']} {fila['apellidoPaterno']}"

    page.locator(SELECTOR_BOTON_CREAR_NUEVO).first.click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.3)

    rut_input = page.locator(SELECTOR_RUT_CREAR)
    rut_input.click()
    rut_input.fill(rut)
    rut_input.press("Tab")
    time.sleep(0.3)
    page.locator(SELECTOR_BTN_ACEPTAR_RUT).click(timeout=5000)
    page.wait_for_load_state("networkidle")
    try:
        page.wait_for_selector("#txtNombres", timeout=5000)
    except Exception:
        pass
    time.sleep(0.3)

    escribir_campo_texto(page, "txtNombres", str(fila["NOMBRES"]))
    escribir_campo_texto(page, "txtApellidoPaterno", str(fila["apellidoPaterno"]))
    escribir_campo_texto(page, "txtApellidoMaterno", str(fila["apellidoMaterno"]))
    escribir_campo_texto(page, "txtSueldoBase", str(fila["sueldoBase"]))
    escribir_combobox_simple(page, "cbSexo", fila["SEXO"])
    escribir_combobox_simple(page, "cbAFP", fila["AFP"])
    escribir_combobox_simple(page, "cbIsapre", fila["ISAPRE"])
    escribir_campo_texto(page, "calendarioFechaInicio_txtCalendar", formatear_fecha(fila["fechaContratacion"]))
    escribir_campo_texto(page, "calendarioFechaTermino_txtCalendar", formatear_fecha(fila["fechaTermino"]))

    if not establecer_multiselect_valor_unico(page, "Proveedores:", fila["PROVEEDOR"]):
        cerrar_formulario(page)
        return ResultadoFila(rut=rut, nombre_excel=nombre_completo, estado="ERROR",
                              detalle=f"Proveedor '{fila['PROVEEDOR']}' no existe en el catálogo real.")

    confirmar_proveedor_seleccionado(page)

    cargo_deseado = fila["CARGO"]
    if not validar_cargo_existe(page, cargo_deseado):
        cerrar_formulario(page)
        return ResultadoFila(rut=rut, nombre_excel=nombre_completo, estado="OMITIDO_CARGO",
                              detalle=f"Cargo '{cargo_deseado}' no existe en el catálogo real. Creación omitida.")
    establecer_multiselect_valor_unico(page, "Categoría Trabajador:", cargo_deseado)

    if not establecer_multiselect_valor_unico(page, "Tiendas:", fila["TIENDA"]):
        cerrar_formulario(page)
        return ResultadoFila(rut=rut, nombre_excel=nombre_completo, estado="ERROR",
                              detalle=f"Tienda '{fila['TIENDA']}' no existe en el catálogo real.")

    page.locator(SELECTOR_BTN_GUARDAR).click(timeout=5000)
    page.wait_for_load_state("networkidle")
    time.sleep(0.5)

    return ResultadoFila(rut=rut, nombre_excel=nombre_completo, estado="CREADO",
                          detalle="Colaborador creado con los datos del Excel.")


# ---------------------------------------------------------------------------
# REPORTE DE SALIDA
# ---------------------------------------------------------------------------

COLORES_ESTADO = {
    "CREADO": "C6EFCE",
    "EDITADO": "C6EFCE",
    "SIN_CAMBIOS": "D9E1F2",
    "OMITIDO_CARGO": "FFEB9C",
    "ERROR": "FFC7CE",
}


def escribir_reporte(resultados: list[ResultadoFila], output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"

    encabezados = ["RUT", "Nombre (Excel)", "Estado", "Detalle"]
    ws.append(encabezados)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in resultados:
        ws.append([r.rut, r.nombre_excel, r.estado, r.detalle])
        fill = PatternFill(start_color=COLORES_ESTADO.get(r.estado, "FFFFFF"),
                            end_color=COLORES_ESTADO.get(r.estado, "FFFFFF"),
                            fill_type="solid")
        for cell in ws[ws.max_row]:
            cell.fill = fill

    for col_cells in ws.columns:
        largo = max(len(str(c.value)) if c.value else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(largo + 4, 80)

    try:
        wb.save(output_path)
    except PermissionError:
        alterno = f"{output_path.rsplit('.', 1)[0]}_{int(time.time())}.xlsx"
        print(f"\nADVERTENCIA: no se pudo guardar en '{output_path}' "
              f"(¿está abierto en Excel u otro programa?). Guardando como '{alterno}' en su lugar.")
        wb.save(alterno)
        output_path = alterno

    print(f"\nReporte guardado en: {output_path}")


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Bot de creación/edición — ShiftLaboral (Fase 2)")
    parser.add_argument("--input", required=True, help="Ruta al Excel de colaboradores a crear/editar")
    parser.add_argument("--output", default="reporte_crear.xlsx", help="Ruta del Excel de salida")
    args = parser.parse_args()

    df = cargar_excel(args.input)
    print(f"Cargados {len(df)} colaboradores desde {args.input}")

    playwright, browser, page = conectar_a_chrome_existente()
    asegurar_pagina_trabajadores(page)

    resultados: list[ResultadoFila] = []
    grupo_actual = None

    for i, fila in df.iterrows():
        rut = str(fila["RUT"]).strip()
        nombre_completo = f"{fila['NOMBRES']} {fila['apellidoPaterno']}"
        print(f"[{i+1}/{len(df)}] Procesando RUT {rut} ({nombre_completo})...")

        try:
            proveedor = str(fila["PROVEEDOR"]).strip()
            if proveedor != grupo_actual:
                seleccionar_grupo_proveedor(page, proveedor)
                grupo_actual = proveedor

            encontrado = buscar_rut(page, rut)

            if encontrado:
                resultado = editar_trabajador_existente(page, fila)
            else:
                resultado = crear_trabajador_nuevo(page, fila)

            resultados.append(resultado)
            print(f"   -> {resultado.estado}: {resultado.detalle}")

            if resultado.estado in ("CREADO", "EDITADO"):
                # Pedido explícito del usuario: tras un guardado exitoso,
                # resetear la vista pasando por el menú antes de seguir, para
                # no arrastrar ningún estado residual del guardado anterior.
                resetear_vista_trabajadores(page)
                grupo_actual = None  # se perdió al navegar, hay que reseleccionarlo
            else:
                limpiar_filtro(page)

        except Exception as e:
            resultados.append(ResultadoFila(rut=rut, nombre_excel=nombre_completo,
                                             estado="ERROR",
                                             detalle=f"Error inesperado durante el procesamiento: {e}"))
            print(f"   -> ERROR: {e}")
            try:
                cerrar_formulario(page)
                asegurar_pagina_trabajadores(page)
            except Exception:
                pass

    escribir_reporte(resultados, args.output)

    total = len(resultados)
    creados = sum(1 for r in resultados if r.estado == "CREADO")
    editados = sum(1 for r in resultados if r.estado == "EDITADO")
    sin_cambios = sum(1 for r in resultados if r.estado == "SIN_CAMBIOS")
    omitidos = sum(1 for r in resultados if r.estado == "OMITIDO_CARGO")
    errores = sum(1 for r in resultados if r.estado == "ERROR")
    print(f"\nResumen: {total} procesados | Creados: {creados} | Editados: {editados} | "
          f"Sin cambios: {sin_cambios} | Omitidos (Cargo): {omitidos} | Error: {errores}")

    browser.close()
    playwright.stop()


if __name__ == "__main__":
    main()
